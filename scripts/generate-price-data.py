"""Dump the Excel price list into src/data/priceData.json.

Rows keep their original structure (category, product, mm/inch sizes, spec,
packing qty, rate, note) so the frontend can group them into product variant
tables. Committed to the repo — the site never needs Excel parsing at runtime.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "jp-plastic-price-list-2082-09-01.xlsx"
OUT = ROOT / "src" / "data" / "priceData.json"


def num(v) -> float | None:
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Price List"]
    rows = []
    for r in range(2, ws.max_row + 1):
        cat = ws.cell(row=r, column=1).value
        prod = ws.cell(row=r, column=2).value
        if not cat or not prod:
            continue
        rows.append(
            {
                "category": str(cat).strip(),
                "product": str(prod).strip(),
                "mm": (str(ws.cell(row=r, column=3).value).strip() if ws.cell(row=r, column=3).value else None),
                "inch": (str(ws.cell(row=r, column=4).value).strip() if ws.cell(row=r, column=4).value else None),
                "spec": (str(ws.cell(row=r, column=5).value).strip() if ws.cell(row=r, column=5).value else None),
                "packing": num(ws.cell(row=r, column=6).value),
                "rate": num(ws.cell(row=r, column=7).value),
                "note": (str(ws.cell(row=r, column=8).value).strip() if ws.cell(row=r, column=8).value else None),
            }
        )
    meta = {
        "effective": "2082/09/01",
        "source": "jp-plastic-price-list-2082-09-01.xlsx",
        "hdpeExcludesVat": True,
        "tankRateUnit": "Rs per litre",
        "notes": [
            "Prices are based at Ex-factory Chitwan, inclusive of 13% VAT and Excise duty — except the HDPE price list.",
            "Prices are subject to change without any prior notice.",
            "This price list replaces all previous price lists.",
        ],
    }
    OUT.write_text(json.dumps({"meta": meta, "rows": rows}, indent=1))
    print(f"wrote {OUT} — {len(rows)} rows")
    cats: dict[str, int] = {}
    for row in rows:
        cats[row["category"]] = cats.get(row["category"], 0) + 1
    for c, n in cats.items():
        print(f"  {n:4d}  {c}")


if __name__ == "__main__":
    main()
