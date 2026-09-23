#!/usr/bin/env python3
"""Dump every sheet of the price-list xlsx: merged ranges, dimensions, all cells."""
import openpyxl

wb = openpyxl.load_workbook("jp-plastic-price-list-2082-09-01.xlsx", data_only=True)
print("SHEETS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n{'='*80}\nSHEET: {name}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("MERGED:", [str(r) for r in ws.merged_cells.ranges][:40])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        cells = []
        for c in row:
            if c.value is not None:
                v = str(c.value).replace("\n", "\\n")
                cells.append(f"{c.coordinate}={v!r}")
        if cells:
            print(" | ".join(cells))
